"""
Tests for new tools: PPT, Excel, Knowledge Base
"""
import sys
import os
import asyncio
from pathlib import Path
from unittest.mock import patch, MagicMock

sys.path.insert(0, str(Path(__file__).parent.parent / "backend"))

import pytest

SAMPLE_DATA = {
    "summary": "Quarterly electrical inspection identified 3 high-severity issues requiring immediate action.",
    "findings": [
        "Main 11kV switchgear SG-4A — thermal discoloration, temperature 78°C vs rated 55°C",
        "MCC-4B contactor for pump P-401A — severe pitting, estimated life < 2 weeks",
        "Earthing resistance at TR-4C = 14.8 ohms (limit: 5 ohms)",
    ],
    "risks": [
        "Risk of arc flash at switchgear SG-4A — HIGH severity",
        "Production loss 450 MT/day if crude feed pump trips — HIGH",
        "Personnel electrocution risk from earthing failure — HIGH",
    ],
    "recommendations": [
        "De-energize SG-4A Bus Section B within 24 hours, arrange OEM inspection",
        "Replace MCC-4B contactor for P-401A within 48 hours",
        "Commission earthing remediation at TR-4C within 7 days",
    ],
}


# ─────────────────────────────────────────────────────────────────────────────
#  PowerPoint Tool Tests
# ─────────────────────────────────────────────────────────────────────────────

class TestDraftPPT:

    @pytest.mark.asyncio
    async def test_generates_pptx_file(self):
        try:
            from pptx import Presentation
        except ImportError:
            pytest.skip("python-pptx not installed")

        from tools.draft_ppt import run_draft_ppt

        result = await run_draft_ppt(
            extracted_data=SAMPLE_DATA,
            title="Electrical Inspection — Board Briefing",
            subtitle="Unit 4 — CDU Quarterly Report",
            document_type="board_presentation",
        )

        assert result["success"] is True, f"PPT generation failed: {result.get('error')}"
        assert result["file_path"].endswith(".pptx")
        assert os.path.exists(result["file_path"])
        assert result["slides"] >= 5  # title + agenda + summary + findings + risks + recs + close

        # Verify it's a valid pptx
        prs = Presentation(result["file_path"])
        assert len(prs.slides) >= 5

        print(f"\n✅ PPT: {result['filename']} ({result['slides']} slides)")
        os.unlink(result["file_path"])

    @pytest.mark.asyncio
    async def test_handles_empty_data(self):
        try:
            from pptx import Presentation
        except ImportError:
            pytest.skip("python-pptx not installed")

        from tools.draft_ppt import run_draft_ppt

        result = await run_draft_ppt({}, title="Empty Test", document_type="test")
        # Should still produce title + closing at minimum
        assert result["success"] is True
        assert os.path.exists(result["file_path"])
        os.unlink(result["file_path"])

    @pytest.mark.asyncio
    async def test_result_has_required_keys(self):
        try:
            from pptx import Presentation
        except ImportError:
            pytest.skip("python-pptx not installed")

        from tools.draft_ppt import run_draft_ppt
        result = await run_draft_ppt(SAMPLE_DATA)
        for key in ["success", "file_path", "filename", "slides", "error"]:
            assert key in result
        if result["success"]:
            os.unlink(result["file_path"])


# ─────────────────────────────────────────────────────────────────────────────
#  Excel Tool Tests
# ─────────────────────────────────────────────────────────────────────────────

class TestDraftExcel:

    @pytest.mark.asyncio
    async def test_generates_xlsx_file(self):
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        from tools.draft_excel import run_draft_excel

        result = await run_draft_excel(
            extracted_data=SAMPLE_DATA,
            title="CDU Electrical Inspection Data",
            document_type="inspection_data",
        )

        assert result["success"] is True, f"Excel failed: {result.get('error')}"
        assert result["file_path"].endswith(".xlsx")
        assert os.path.exists(result["file_path"])
        assert result["sheets"] >= 3  # Cover + Findings + Risk Register + Action Plan

        # Verify content
        wb = openpyxl.load_workbook(result["file_path"])
        assert "Findings" in wb.sheetnames
        assert "Risk Register" in wb.sheetnames
        assert "Action Plan" in wb.sheetnames

        findings_ws = wb["Findings"]
        # Check header row
        assert findings_ws.cell(2, 2).value == "Finding Description"

        print(f"\n✅ Excel: {result['filename']} ({result['sheets']} sheets)")
        os.unlink(result["file_path"])

    @pytest.mark.asyncio
    async def test_findings_populated(self):
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        from tools.draft_excel import run_draft_excel
        result = await run_draft_excel(SAMPLE_DATA)
        assert result["success"] is True
        wb = openpyxl.load_workbook(result["file_path"])
        ws = wb["Findings"]
        # Row 3 should be first finding
        assert ws.cell(3, 2).value == SAMPLE_DATA["findings"][0]
        os.unlink(result["file_path"])

    @pytest.mark.asyncio
    async def test_result_has_required_keys(self):
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        from tools.draft_excel import run_draft_excel
        result = await run_draft_excel(SAMPLE_DATA)
        for key in ["success", "file_path", "filename", "sheets", "error"]:
            assert key in result
        if result["success"]:
            os.unlink(result["file_path"])


# ─────────────────────────────────────────────────────────────────────────────
#  Knowledge Base Tests
# ─────────────────────────────────────────────────────────────────────────────

class TestKnowledgeBase:

    @pytest.fixture(autouse=True)
    def use_temp_kb(self, tmp_path, monkeypatch):
        """Redirect KB storage to a temp dir for test isolation."""
        import tools.knowledge_base as kb_module
        monkeypatch.setattr(kb_module, "KB_DIR",       str(tmp_path))
        monkeypatch.setattr(kb_module, "KB_INDEX_FILE", str(tmp_path / "index.pkl"))
        monkeypatch.setattr(kb_module, "KB_DOCS_FILE",  str(tmp_path / "documents.json"))

    def test_ingest_text(self):
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError:
            pytest.skip("sentence-transformers not installed")

        from tools.knowledge_base import ingest_text

        result = ingest_text(
            text="Earthing resistance must not exceed 5 ohms per IS:3043. "
                 "Earth pits must be tested semi-annually using fall-of-potential method.",
            source_name="SOP-ELE-012",
            doc_type="manual",
        )
        assert result["success"] is True
        assert result["chunks_added"] >= 1
        assert result["source"] == "SOP-ELE-012"

    def test_search_returns_results(self):
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError:
            pytest.skip("sentence-transformers not installed")

        from tools.knowledge_base import ingest_text, search_knowledge_base

        # Ingest some content
        ingest_text(
            "Earthing resistance limit is 5 ohms. Test with fall-of-potential method.",
            "SOP-ELE-012", "manual"
        )
        ingest_text(
            "Emergency lighting must provide 10 lux at floor level per IS:3646.",
            "SOP-SAFETY-003", "manual"
        )
        ingest_text(
            "Zone 1 conduit sealing fittings must be resealed quarterly.",
            "SOP-ELE-012", "manual"
        )

        result = search_knowledge_base("earthing resistance standard limit")
        assert result["success"] is True
        assert len(result["results"]) >= 1
        # Top result should be about earthing
        top = result["results"][0]
        assert "earthing" in top["text"].lower() or "earth" in top["text"].lower()
        assert top["score"] > 0.15
        assert top["source"] == "SOP-ELE-012"

    def test_search_empty_kb_returns_error(self):
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError:
            pytest.skip("sentence-transformers not installed")

        from tools.knowledge_base import search_knowledge_base

        result = search_knowledge_base("any query")
        assert result["success"] is False
        assert "empty" in result["error"].lower()

    def test_get_kb_stats_empty(self):
        from tools.knowledge_base import get_kb_stats
        stats = get_kb_stats()
        assert "total_chunks" in stats
        assert "sources" in stats
        assert "ready" in stats
        assert stats["ready"] is False

    def test_get_kb_stats_after_ingest(self):
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError:
            pytest.skip("sentence-transformers not installed")

        from tools.knowledge_base import ingest_text, get_kb_stats

        ingest_text("Test content about electrical safety.", "test_doc", "manual")
        stats = get_kb_stats()
        assert stats["total_chunks"] >= 1
        assert "test_doc" in stats["sources"]
        assert stats["ready"] is True

    def test_no_duplicate_ingestion(self):
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError:
            pytest.skip("sentence-transformers not installed")

        from tools.knowledge_base import ingest_text

        text = "Same text ingested twice to test deduplication."
        r1 = ingest_text(text, "source_a", "manual")
        r2 = ingest_text(text, "source_a", "manual")
        assert r1["chunks_added"] >= 1
        assert r2["chunks_added"] == 0  # duplicate should be skipped

    @pytest.mark.asyncio
    async def test_ingest_txt_file(self, tmp_path):
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError:
            pytest.skip("sentence-transformers not installed")

        from tools.knowledge_base import run_ingest_file

        # Create a temp text file
        txt_file = tmp_path / "test_doc.txt"
        txt_file.write_text(
            "Earthing systems must comply with IS:3043.\n"
            "Quarterly tests required for all HT equipment.\n"
            "Maximum earthing resistance is 5 ohms.",
            encoding="utf-8"
        )
        result = await run_ingest_file(str(txt_file), source_name="test_doc.txt", doc_type="manual")
        assert result["success"] is True
        assert result["chunks_added"] >= 1

    @pytest.mark.asyncio
    async def test_ingest_nonexistent_file(self):
        from tools.knowledge_base import run_ingest_file
        result = await run_ingest_file("/nonexistent/file.txt")
        assert result["success"] is False
        assert "not found" in result["error"].lower()

    @pytest.mark.asyncio
    async def test_async_search(self):
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError:
            pytest.skip("sentence-transformers not installed")

        from tools.knowledge_base import ingest_text, run_search_kb

        ingest_text("Zone 1 conduit seals must be inspected quarterly.", "SOP", "manual")
        result = await run_search_kb("conduit seal inspection")
        assert result["success"] is True


# ─────────────────────────────────────────────────────────────────────────────
#  Fixture File Tests (verify demo fixtures exist and are valid)
# ─────────────────────────────────────────────────────────────────────────────

class TestFixtures:
    FIXTURE_DIR = Path(__file__).parent / "fixtures"

    def test_fixtures_directory_exists(self):
        assert self.FIXTURE_DIR.exists(), "Run: python tests/fixtures/generate_fixtures.py"

    def test_sop_fixture_exists(self):
        sop = self.FIXTURE_DIR / "sample_sop.txt"
        if not sop.exists():
            pytest.skip("Fixtures not generated yet — run generate_fixtures.py")
        text = sop.read_text(encoding="utf-8")
        assert "MRPL" in text
        assert "earthing" in text.lower()
        assert len(text) > 1000

    def test_board_brief_fixture_exists(self):
        brief = self.FIXTURE_DIR / "sample_board_brief.txt"
        if not brief.exists():
            pytest.skip("Fixtures not generated yet")
        text = brief.read_text(encoding="utf-8")
        assert "INR" in text
        assert "switchgear" in text.lower()

    def test_code_problem_fixture_exists(self):
        code = self.FIXTURE_DIR / "sample_code_problem.py"
        if not code.exists():
            pytest.skip("Fixtures not generated yet")
        text = code.read_text(encoding="utf-8")
        assert "ASME" in text
        assert "def calculate_pipe_wall_thickness" in text
        # The sandbox demo runs this file, so it must at least be valid Python
        compile(text, str(code), "exec")

    def test_pid_diagram_exists(self):
        pid = self.FIXTURE_DIR / "sample_pid_diagram.png"
        if not pid.exists():
            pytest.skip("Fixtures not generated yet")
        from PIL import Image
        img = Image.open(pid)
        assert img.size[0] >= 1000  # should be at least 1000px wide
        assert img.size[1] >= 800

    def test_inspection_report_exists(self):
        report = self.FIXTURE_DIR / "sample_inspection_report.pdf"
        txt_fallback = self.FIXTURE_DIR / "sample_inspection_report.txt"
        if not report.exists() and not txt_fallback.exists():
            pytest.skip("Fixtures not generated yet")
        # At least one should exist
        assert report.exists() or txt_fallback.exists()
